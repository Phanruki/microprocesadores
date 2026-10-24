#!/usr/bin/env python3
"""
Genera la presentacion de la exposicion a partir de la plantilla UV.

    plantilla : docs/practica/presentacion_plantilla.pptx   (obligatoria;
                si no esta, se usa docs/practica/final/plantillas/...)
    salida    : docs/reporte_exposicion/exposicion/presentacion_exposicion.pptx

Uso (desde la raiz del proyecto):
    python3 docs/reporte_exposicion/exposicion/generar_presentacion.py

La presentacion es de APOYO: pocas palabras por diapositiva, la explicacion
la da el equipo. El guion de lo que hay que decir esta en
docs/reporte_exposicion/exposicion/guion_exposicion.md
"""

import os

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

RAIZ = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)),
                                     "..", "..", ".."))

# La plantilla obligatoria puede estar copiada en docs/practica/ o en su
# ubicacion original del repositorio; se usa la primera que exista.
_PLANTILLAS = (
    os.path.join(RAIZ, "docs", "practica", "presentacion_plantilla.pptx"),
    os.path.join(RAIZ, "docs", "practica", "final", "plantillas",
                 "presentacion_plantilla.pptx"),
)
PLANTILLA = next((p for p in _PLANTILLAS if os.path.exists(p)), _PLANTILLAS[-1])

SALIDA = os.path.join(RAIZ, "docs", "reporte_exposicion", "exposicion",
                      "presentacion_exposicion.pptx")
IMG = os.path.join(RAIZ, "docs", "reporte_exposicion", "reporte_final", "imagenes")
DIAG = os.path.join(RAIZ, "docs", "diagramas")

# ---------------------------------------------------------------------------
# Paleta
# ---------------------------------------------------------------------------
AZUL = RGBColor(0x1F, 0x49, 0x7D)        # azul UV / titulos
AZUL_MED = RGBColor(0x4F, 0x81, 0xBD)    # azul intermedio
AZUL_CLARO = RGBColor(0xDC, 0xE6, 0xF1)  # relleno claro
VERDE = RGBColor(0x2E, 0x7D, 0x32)
ROJO = RGBColor(0xC0, 0x50, 0x4D)
GRIS = RGBColor(0x40, 0x40, 0x40)
GRIS_CLARO = RGBColor(0xF3, 0xF6, 0xFA)
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)
NEGRO = RGBColor(0x1A, 0x1A, 0x1A)

CODIGO = "Consolas"
TITULO = "Gill Sans MT"

# Geometria de los layouts (en pulgadas) -----------------------------------
# layout 1 -> titulo (0.50, 0.75)  cuerpo (1.344, 1.301) 8.155 x 3.713
# layout 2 -> titulo (1.314, 2.264) 8.155 x 0.469 ; cuerpo (1.314, 2.695)
# layout 3 -> titulos yuerq/der   contenido yuerq/der (dos columnas)
# layout 4 -> titulo (0.484, 0.989) ; texto izq (0.484, 1.614) 3.297 x 3.387
#                     contenido der (4.068, 1.614) 5.305 x 3.386
# layout 5 -> titulo (0.675, 4.284) ; imagen (1.753, 1.019) 6.494 x 3.192
#                     pie (0.669, 4.702)

EMU_IN = 914400


def IN(valor):
    return Emu(int(round(valor * EMU_IN)))


def ancho(shape):
    return shape.width / EMU_IN


def alto(shape):
    return shape.height / EMU_IN


# ---------------------------------------------------------------------------
# Helpers de texto
# ---------------------------------------------------------------------------
def _runs(p, texto, size, color, font=None, bold=False, italic=False):
    """Escribe un texto admitiendo **negritas**."""
    for i, parte in enumerate(texto.split("**")):
        if parte == "":
            continue
        r = p.add_run()
        r.text = parte
        f = r.font
        f.size = Pt(size)
        f.color.rgb = color
        f.bold = bold or (i % 2 == 1)
        f.italic = italic
        if font:
            f.name = font


def _pconf(p, marL=None, indent=None, space_after=None, space_before=None,
           line=None, align=None):
    if align is not None:
        p.alignment = align
    if space_after is not None:
        p.space_after = Pt(space_after)
    if space_before is not None:
        p.space_before = Pt(space_before)
    if line is not None:
        p.line_spacing = line
    if marL is not None or indent is not None:
        pPr = p._p.get_or_add_pPr()
        if marL is not None:
            pPr.set("marL", str(int(round(marL * EMU_IN))))
        if indent is not None:
            pPr.set("indent", str(int(round(indent * EMU_IN))))


def nuevo(prs, layout_idx):
    return prs.slides.add_slide(prs.slide_layouts[layout_idx])


def borrar_ph(slide, idx):
    try:
        ph = slide.placeholders[idx]
    except KeyError:
        return
    el = ph._element
    el.getparent().remove(el)


def titulo(slide, texto, size=24, algn=PP_ALIGN.LEFT, color=AZUL, bold=True, idx=None):
    ph = slide.placeholders[idx] if idx is not None else slide.shapes.title
    tf = ph.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    _pconf(p, align=algn)
    _runs(p, texto, size, color, font=TITULO, bold=bold)
    return ph


def lista(slide, idx, items, size=15, color=GRIS, marco="•", gap=9,
          line=1.15, align=PP_ALIGN.LEFT):
    """items: cadena (nivel 0) o (nivel, cadena). Admite **negritas**."""
    ph = slide.placeholders[idx]
    tf = ph.text_frame
    tf.clear()
    tf.word_wrap = True
    primero = True
    for item in items:
        nivel, texto = (0, item) if isinstance(item, str) else item
        p = tf.paragraphs[0] if primero else tf.add_paragraph()
        primero = False
        sangria = 0.26 + 0.22 * nivel
        _pconf(p, marL=sangria, indent=-sangria, space_after=gap, line=line, align=align)
        if texto:
            guion = ("   " * nivel) + (marco if nivel == 0 else "–  ")
            _runs(p, guion + texto, size if nivel == 0 else size - 1.5, color)
    return ph


def bloque(slide, x, y, w, h, lineas, fill=BLANCO, borde=AZUL, size=10,
           color=GRIS, align=PP_ALIGN.LEFT, cabecera=None, font=None,
           shape=MSO_SHAPE.ROUNDED_RECTANGLE, radio=0.08, anchor=MSO_ANCHOR.TOP,
           marco=None, gap=3):
    """lineas: lista de textos (una por parrafo)."""
    sh = slide.shapes.add_shape(shape, IN(x), IN(y), IN(w), IN(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if borde is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = borde
        sh.line.width = Pt(1)
    try:
        sh.adjustments[0] = radio
    except Exception:
        pass
    sh.shadow.inherit = False
    tf = sh.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = IN(0.07)
    tf.margin_right = IN(0.07)
    tf.margin_top = IN(0.04)
    tf.margin_bottom = IN(0.04)
    primera = True
    if cabecera is not None:
        p = tf.paragraphs[0]
        _pconf(p, space_after=gap, align=PP_ALIGN.CENTER)
        _runs(p, cabecera, size + 1.5, BLANCO, bold=True)
        primera = False
    for texto in lineas:
        p = tf.paragraphs[0] if primera else tf.add_paragraph()
        primera = False
        _pconf(p, space_after=gap, line=1.0, align=align)
        if marco:
            _pconf(p, marL=0.2, indent=-0.2)
            texto = marco + texto
        _runs(p, texto, size, color, font=font)
    return sh


def etiqueta(slide, x, y, w, h, texto, size=12, color=AZUL, bold=True,
             align=PP_ALIGN.CENTER, font=None, fill=None, borde=None):
    tb = slide.shapes.add_textbox(IN(x), IN(y), IN(w), IN(h))
    tf = tb.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    if fill is not None:
        tb.fill.solid()
        tb.fill.fore_color.rgb = fill
    if borde is not None:
        tb.line.color.rgb = borde
    p = tf.paragraphs[0]
    _pconf(p, align=align, line=1.05)
    _runs(p, texto, size, color, font=font, bold=bold)
    return tb


def caja(slide, x, y, w, h, texto, size=10, fill=AZUL, color=BLANCO, bold=True,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, borde=None, align=PP_ALIGN.CENTER,
         font=None):
    sh = slide.shapes.add_shape(shape, IN(x), IN(y), IN(w), IN(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if borde is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = borde
        sh.line.width = Pt(1)
    sh.shadow.inherit = False
    tf = sh.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = IN(0.04)
    tf.margin_right = IN(0.04)
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    _pconf(p, align=align, line=1.0)
    _runs(p, texto, size, color, bold=bold, font=font)
    return sh


def flecha(slide, x1, y1, x2, y2, color=AZUL_MED, width=1.5, dash=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, IN(x1), IN(y1), IN(x2), IN(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    if dash:
        el = ln.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}prstDash",
                            {"val": "dash"})
        ln.append(el)
    tail = ln.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd",
                          {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)
    return c


def tabla(slide, datos, x, y, w, h, col_w=None, size=11, cabecera=True,
          align=PP_ALIGN.LEFT):
    gf = slide.shapes.add_table(len(datos), len(datos[0]), IN(x), IN(y), IN(w), IN(h))
    t = gf.table
    if col_w:
        for i, cw in enumerate(col_w):
            t.columns[i].width = IN(cw)
    for r, fila in enumerate(datos):
        for c, valor in enumerate(fila):
            cell = t.cell(r, c)
            cell.margin_left = IN(0.07)
            cell.margin_right = IN(0.06)
            cell.margin_top = IN(0.03)
            cell.margin_bottom = IN(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.clear()
            tf.word_wrap = True
            p = tf.paragraphs[0]
            _pconf(p, align=align, line=1.0)
            es_cab = cabecera and r == 0
            _runs(p, valor, size, BLANCO if es_cab else GRIS, bold=es_cab)
            cell.fill.solid()
            if es_cab:
                cell.fill.fore_color.rgb = AZUL
            else:
                cell.fill.fore_color.rgb = GRIS_CLARO if r % 2 == 0 else BLANCO
    return t


def imagen(slide, ruta, x, y, w, h):
    """Coloca la imagen dentro del rectangulo conservando la proporcion."""
    iw, ih = Image.open(ruta).size
    esc = min(w / iw, h / ih)
    nw, nh = iw * esc, ih * esc
    return slide.shapes.add_picture(ruta, IN(x + (w - nw) / 2.0),
                                    IN(y + (h - nh) / 2.0), IN(nw), IN(nh))


def imagen_del_ph(slide, idx, ruta):
    ph = slide.placeholders[idx]
    x, y = ph.left / EMU_IN, ph.top / EMU_IN
    w, h = ph.width / EMU_IN, ph.height / EMU_IN
    el = ph._element
    el.getparent().remove(el)
    return imagen(slide, ruta, x, y, w, h)


def separador(prs, numero, texto, pie, layout=2):
    s = nuevo(prs, layout)
    titulo(s, texto, size=26, algn=PP_ALIGN.CENTER, idx=0)
    lista(s, 10, [pie], size=13, color=AZUL_MED, marco="", align=PP_ALIGN.CENTER, gap=0)
    etiqueta(s, 3.9, 1.05, 2.2, 1.0, numero, size=54, color=AZUL_CLARO,
             align=PP_ALIGN.CENTER)
    return s


# ---------------------------------------------------------------------------
# Construccion
# ---------------------------------------------------------------------------
def main():
    prs = Presentation(PLANTILLA)

    # -----------------------------------------------------------------
    # 1. Portada (diapositiva existente de la plantilla)
    # -----------------------------------------------------------------
    portada = prs.slides[0]
    for sh in portada.shapes:
        if sh.name == "CuadroTexto 4":
            tf = sh.text_frame
            tf.clear()
            lineas = [
                ("Microprocesadores y Microcontroladores", 9, False),
                ("", 9, False),
                ("Emulador del procesador MIPS en lenguaje Java", 12, True),
                ("", 9, False),
                ("Presenta:", 9, False),
                ("Soto Hernandez Yadid Asiel (zS24006887)", 9, True),
                ("Berman Islas Jose de Jesus (zS24006897)", 9, True),
                ("Damian Bumas Miguel Alejandro (zS24021396)", 9, True),
                ("", 9, False),
                ("23 de Octubre de 2026", 9, False),
                ("“Lis de Veracruz: Arte, Ciencia, Luz”", 9, False),
            ]
            for i, (txt, sz, bold) in enumerate(lineas):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                _pconf(p, align=PP_ALIGN.RIGHT, space_after=0, line=1.05)
                if txt:
                    r = p.add_run()
                    r.text = txt
                    r.font.size = Pt(sz)
                    r.font.bold = bold
                    r.font.name = "Gill Sans MT"
            break

    # -----------------------------------------------------------------
    # 2. Contenido (diapositiva existente de la plantilla)
    # -----------------------------------------------------------------
    contenido = prs.slides[1]
    ph = contenido.placeholders[10]
    ph.left, ph.top = IN(1.344), IN(1.30)
    ph.width, ph.height = IN(8.155), IN(3.713)
    puntos = [
        "Razones tecnicas para elegir el procesador",
        "Arquitectura: PC, ROM, RAM y registros",
        "ISA del procesador y pruebas tecnicas",
        "Codigos implementados y ejecutados",
    ]
    tf = ph.text_frame
    parrafos = list(tf.paragraphs)
    for i, txt in enumerate(puntos):
        p = parrafos[i]
        if p.runs:
            p.runs[0].text = txt
            for extra in p.runs[1:]:
                extra._r.getparent().remove(extra._r)
    for p in parrafos[len(puntos):]:
        p._p.getparent().remove(p._p)

    # =================================================================
    # BLOQUE 1 - Razones tecnicas (min 1-5)
    # =================================================================
    separador(prs, "01", "Razones tecnicas para elegir el procesador",
              "Por que MIPS  ·  por que un emulador          min. 1 – 5")

    # --- Que construimos -------------------------------------------------
    s = nuevo(prs, 1)
    titulo(s, "¿Que construimos?")
    lista(s, 10, [
        "Emulador de un **procesador MIPS** de 32 bits, escrito en **Java**",
        "Sin librerias externas: **36 clases**, 22 de ellas en el nucleo",
        "5 paquetes: cpu · ejecucion · memoria · alu · instruccion",
        "Ciclo completo de **5 etapas**: IF -> ID -> EX -> MEM -> WB",
        "**5 programas** de usuario y **201 pruebas** automaticas",
        "Objetivo: ver por dentro lo que normalmente es una caja negra",
    ], size=17, gap=13)

    # --- Emulador vs procesador real ------------------------------------
    s = nuevo(prs, 3)
    titulo(s, "Procesador real", size=19, idx=17, color=ROJO)
    titulo(s, "Emulador (el nuestro)", size=19, idx=16, color=VERDE)
    lista(s, 14, [
        "Millones de transistores",
        "Diseno propiedad industrial y secreto",
        "Un error cuesta **silicio** perdido",
        "No se puede mirar su interior",
        "Costo y equipo fuera de alcance",
    ], size=14, gap=9)
    lista(s, 15, [
        "Se construye entero en software",
        "Registros, memoria y senales **a la vista**",
        "Un error es una **prueba fallida** que se corrige",
        "Se observa el ciclo **instruccion por instruccion**",
        "Mismos conceptos que el hardware real",
    ], size=14, gap=9)

    # --- MIPS frente a las alternativas ---------------------------------
    s = nuevo(prs, 1)
    titulo(s, "MIPS frente a las alternativas")
    borrar_ph(s, 10)
    tabla(s, [
        ["Criterio", "MIPS (elegido)", "RISC-V", "x86"],
        ["Familia", "RISC", "RISC", "CISC"],
        ["Tamano de instruccion", "32 bits, longitud fija", "32 bits, longitud fija",
         "1 a 15 bytes, variable"],
        ["Formatos", "3: R, I y J", "6", "decenas"],
        ["Especificacion", "publica y documentacion clara", "abierta y modular",
         "propietaria"],
        ["¿Se implementa en un curso?", "si, decodificador simple", "si", "no es practico"],
    ], 0.55, 1.30, 8.9, 2.75, col_w=[2.3, 2.3, 2.15, 2.15], size=10.5)
    etiqueta(s, 0.55, 4.42, 8.9, 0.5,
             "Tambien valia RISC-V; MIPS gano por su documentacion didactica. "
             "x86 (longitud variable + microcodigo) estaba fuera de alcance.",
             size=11, color=AZUL_MED, bold=False)

    # --- Tres razones tecnicas ------------------------------------------
    s = nuevo(prs, 1)
    titulo(s, "Tres razones tecnicas")
    lista(s, 10, [
        "**1. Es RISC**  ->  conjunto reducido, todas las instrucciones miden 32 bits",
        "**2. Tres formatos fijos (R, I, J)**  ->  decodificar = máscaras + desplazamientos",
        "**3. Documentacion abierta y didactica**  ->  mismos conceptos que en cualquier CPU",
        "",
        "**Contra x86:** instrucciones de longitud variable y microcodigo; aqui el "
        "decodificador cabria en una pantalla",
        "**Ganancia:** se implementa el ciclo completo sin perderse en el formato",
    ], size=16, gap=12)

    # =================================================================
    # BLOQUE 2 - Arquitectura (min 6-15)
    # =================================================================
    separador(prs, "02", "Arquitectura del emulador",
              "PC · ROM · RAM · registros · memorias de instrucciones y de datos     "
              "min. 6 – 15")

    # --- 5 paquetes ------------------------------------------------------
    s = nuevo(prs, 1)
    titulo(s, "El sistema en 5 paquetes")
    borrar_ph(s, 10)
    paquetes = [
        ("cpu", "CPU\nPC\nciclo de 5 etapas"),
        ("ejecucion", "UnidadControl\nEjecutorInstrucciones\nSenalesControl"),
        ("memoria", "MemoriaSistema\nROM · RAM · DiscoDuro\nRegistros"),
        ("alu", "ALU\nOperacionALU"),
        ("instruccion", "Instruccion · Decodificador\nOpcode · Funct\nProgramaBuilder"),
    ]
    x0, w, gap, ytop = 0.55, 1.70, 0.10, 1.35
    for i, (nom, clases) in enumerate(paquetes):
        x = x0 + i * (w + gap)
        bloque(s, x, ytop, w, 0.42, [], fill=AZUL, borde=None, size=1)
        etiqueta(s, x, ytop + 0.09, w, 0.3, nom, size=12, color=BLANCO, font=CODIGO)
        bloque(s, x, ytop + 0.44, w, 1.35, clases.split("\n"), fill=GRIS_CLARO,
               borde=AZUL, size=9.5, color=GRIS, align=PP_ALIGN.CENTER, gap=4)
        if i < len(paquetes) - 1:
            flecha(s, x + w + 0.01, ytop + 1.1, x + w + gap - 0.01, ytop + 1.1)
    bloque(s, 0.55, 3.25, 8.9, 0.55,
           ["**Fuera del nucleo** (usan el nucleo, no lo tocan):   ui (ventana grafica)   ·   "
            "programas (5)   ·   pruebas (201)"],
           fill=AZUL_CLARO, borde=AZUL_MED, size=12, color=AZUL, align=PP_ALIGN.CENTER)
    etiqueta(s, 0.55, 4.0, 8.9, 0.9,
             "Diagrama de clases completo: docs/diagramas/diagrama_clases_nucleo.png  "
             "(22 clases en el nucleo)", size=11, color=GRIS, bold=False,
             align=PP_ALIGN.CENTER)

    # --- Mapa de memoria -------------------------------------------------
    s = nuevo(prs, 1)
    titulo(s, "Mapa de memoria")
    borrar_ph(s, 10)
    segmentos = [
        ("ROM  ·  texto", "0 – 99", "100 celdas x 32 bits", "instrucciones", AZUL),
        ("RAM  ·  datos", "100 – 199", "100 celdas x 32 bits", "datos del programa", AZUL_MED),
        ("DiscoDuro  ·  disco", "200 – 299", "100 celdas x 32 bits", "persistencia", GRIS),
    ]
    x0, w, gap, ytop = 0.55, 2.90, 0.10, 1.55
    for i, (nom, rango, celdas, uso, col) in enumerate(segmentos):
        x = x0 + i * (w + gap)
        bloque(s, x, ytop, w, 1.30, [], fill=col, borde=None, size=1)
        etiqueta(s, x + 0.05, ytop + 0.10, w - 0.1, 0.3, nom, size=13, color=BLANCO)
        etiqueta(s, x + 0.05, ytop + 0.45, w - 0.1, 0.3, rango, size=15, color=BLANCO,
                 font=CODIGO)
        etiqueta(s, x + 0.05, ytop + 0.80, w - 0.1, 0.3, celdas, size=10, color=AZUL_CLARO,
                 bold=False)
        etiqueta(s, x + 0.05, ytop + 1.02, w - 0.1, 0.3, uso, size=10, color=AZUL_CLARO,
                 bold=False)
    texto_mapa = [
        "Direccion **global 0 … 299**: una direccion = una **palabra de 32 bits**",
        "**MemoriaSistema** enruta cada acceso al segmento que la contiene "
        "(`segmentoDe`) y avisa si la direccion no existe",
        "El volcado muestra la direccion en **decimal y hexadecimal**: "
        "dir 100 (0x00000064)",
        "**Reglas:** ROM bloqueada tras cargar · RAM arranca en 000 · "
        "DiscoDuro se guarda y se carga en un archivo",
    ]
    tb = s.shapes.add_textbox(IN(0.55), IN(3.35), IN(8.9), IN(1.6))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.clear()
    for i, t in enumerate(texto_mapa):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        _pconf(p, marL=0.24, indent=-0.24, space_after=6, line=1.1)
        _runs(p, "•  " + t, 13, GRIS)

    # --- ROM vs RAM ------------------------------------------------------
    s = nuevo(prs, 3)
    titulo(s, "Memoria de instrucciones (ROM)", size=18, idx=17)
    titulo(s, "Memoria de datos (RAM)", size=18, idx=16)
    lista(s, 14, [
        "Segmento **texto**, direcciones **0 – 99**",
        "100 celdas de **32 bits**",
        "Aqui vive el **programa** (lo que se ensamblo)",
        "**Solo lectura**: `ROM.bloquear()` tras cargar",
        "La CPU busca aqui su instruccion (etapa **IF**)",
    ], size=13, gap=8)
    lista(s, 15, [
        "Segmento **datos**, direcciones **100 – 199**",
        "100 celdas de 32 bits, lectura y escritura",
        "Se accede solo con **lw** y **sw**",
        "Arranca en **000** (verificado por pruebas)",
        "**DiscoDuro** (200 – 299): persistencia en archivo",
    ], size=13, gap=8)

    # --- Registros -------------------------------------------------------
    s = nuevo(prs, 1)
    titulo(s, "Registros del procesador")
    borrar_ph(s, 10)
    nombres = ["$zero", "$at", "$v0", "$v1", "$a0", "$a1", "$a2", "$a3",
               "$t0", "$t1", "$t2", "$t3", "$t4", "$t5", "$t6", "$t7",
               "$s0", "$s1", "$s2", "$s3", "$s4", "$s5", "$s6", "$s7",
               "$t8", "$t9", "$k0", "$k1", "$gp", "$sp", "$fp", "$ra"]
    gw, gh, gg = 0.70, 0.42, 0.06
    x0, y0 = 0.55, 1.45
    for i, nom in enumerate(nombres):
        col, fil = i % 8, i // 8
        x = x0 + col * (gw + gg)
        y = y0 + fil * (gh + gg)
        especial = nom in ("$zero", "$v0", "$a0", "$ra", "$sp")
        caja(s, x, y, gw, gh, nom, size=9.5,
             fill=AZUL if nom == "$zero" else (AZUL_CLARO if especial else GRIS_CLARO),
             color=BLANCO if nom == "$zero" else (AZUL if especial else GRIS),
             borde=AZUL_MED, shape=MSO_SHAPE.RECTANGLE)
    lista_bloque = [
        "**32 registros x 32 bits** dentro de la CPU",
        "`$zero` siempre vale 0: **las escrituras se ignoran**",
        "`$v0` = servicio del syscall · `$a0` = argumento",
        "`$t0–$t9` temporales · `$s0–$s7` salvados",
        "`$ra` retorno · `$sp` pila · `$fp` marco",
        "ISA **load/store**: la aritmetica solo vive en registros",
    ]
    tb = s.shapes.add_textbox(IN(6.75), IN(1.45), IN(2.75), IN(2.9))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.clear()
    for i, t in enumerate(lista_bloque):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        _pconf(p, marL=0.2, indent=-0.2, space_after=7, line=1.1)
        _runs(p, "•  " + t, 10.5, GRIS)
    bloque(s, 0.55, 4.55, 8.9, 0.75,
           ["`Registros.escribir(indice, valor)`  ->  si `indice == 0` la escritura se "
            "descarta; `Registros.leer(0)` devuelve siempre 0"],
           fill=AZUL_CLARO, borde=AZUL, size=12, color=AZUL, align=PP_ALIGN.CENTER,
           font=CODIGO)

    # --- PC ---------------------------------------------------------------
    s = nuevo(prs, 1)
    titulo(s, "PC - Contador de programa")
    borrar_ph(s, 10)
    bx, by, bw, bh = 0.70, 1.75, 0.95, 0.55
    for i in range(5):
        x = bx + i * (bw + 0.45)
        caja(s, x, by, bw, bh, "PC = %d" % i, size=11, fill=AZUL_CLARO, color=AZUL,
             borde=AZUL, shape=MSO_SHAPE.RECTANGLE, font=CODIGO)
        if i < 4:
            flecha(s, x + bw + 0.03, by + bh / 2, x + bw + 0.42, by + bh / 2)
    # salto de vuelta
    x4 = bx + 4 * (bw + 0.45) + bw / 2
    x1 = bx + 1 * (bw + 0.45) + bw / 2
    flecha(s, x4, by + bh + 0.03, x4, by + bh + 0.55, color=ROJO, dash=True)
    flecha(s, x4, by + bh + 0.55, x1, by + bh + 0.55, color=ROJO, dash=True)
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, IN(x1), IN(by + bh + 0.55),
                               IN(x1), IN(by + bh + 0.07))
    c.line.color.rgb = ROJO
    c.line.width = Pt(1.5)
    etiqueta(s, 4.6, by + bh + 0.62, 4.0, 0.3,
             "beq / bne / j / jal / jr escriben otra direccion", size=11, color=ROJO)
    lista_tb = [
        "Registro de **32 bits**: guarda la direccion de la **siguiente** instruccion",
        "Se incrementa en cada busqueda; en nuestro emulador **+1 celda** "
        "(en un MIPS real serian +4 bytes)",
        "Arranca en **0** al encender o reiniciar (`PC.reiniciar()`)",
        "**Circular**: si la direccion se sale del mapa (300) vuelve a 0  ->  "
        "la CPU nunca muere",
        "Codigo: `src/cpu/PC.java` y `CPU.etapaBusqueda()`",
    ]
    tb = s.shapes.add_textbox(IN(0.55), IN(3.45), IN(8.9), IN(1.6))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.clear()
    for i, t in enumerate(lista_tb):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        _pconf(p, marL=0.24, indent=-0.24, space_after=6, line=1.1)
        _runs(p, "•  " + t, 12.5, GRIS)

    # --- 5 etapas ---------------------------------------------------------
    s = nuevo(prs, 1)
    titulo(s, "Ciclo de instruccion: 5 etapas")
    borrar_ph(s, 10)
    etapas = [
        ("IF", "Buscar", "instruccion = ROM[PC]\nPC = PC + 1"),
        ("ID", "Decodificar", "Decodificador separa campos\nUnidad de Control -> senales"),
        ("EX", "Ejecutar", "La ALU calcula\nse resuelven ramas y saltos"),
        ("MEM", "Memoria", "lw lee · sw escribe\nen MemoriaSistema"),
        ("WB", "Escribir", "Llega el resultado\nal registro destino"),
    ]
    cw, step, cy, ch = 2.00, 1.74, 1.55, 0.72
    for i, (sig, nom, desc) in enumerate(etapas):
        x = 0.55 + i * step
        sh = s.shapes.add_shape(MSO_SHAPE.CHEVRON, IN(x), IN(cy), IN(cw), IN(ch))
        sh.fill.solid()
        sh.fill.fore_color.rgb = AZUL if i % 2 == 0 else AZUL_MED
        sh.line.fill.background()
        sh.shadow.inherit = False
        tf = sh.text_frame
        tf.clear()
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        _pconf(p, align=PP_ALIGN.CENTER, line=1.0)
        _runs(p, sig, 15, BLANCO, bold=True)
        p2 = tf.add_paragraph()
        _pconf(p2, align=PP_ALIGN.CENTER, line=1.0)
        _runs(p2, nom, 10, AZUL_CLARO, bold=False)
        etiqueta(s, x + 0.10, cy + ch + 0.15, 1.62, 1.1, desc.replace("\n", "  ·  "),
                 size=10.5, color=GRIS, bold=False, align=PP_ALIGN.LEFT)
    bloque(s, 0.55, 3.75, 8.9, 0.72,
           ["Todo lo que hace un procesador = **repetir este ciclo**: `CPU.ciclo()` "
            "recorre IF -> ID -> EX -> MEM -> WB y suma 1 al contador de ciclos"],
           fill=AZUL_CLARO, borde=AZUL, size=13, color=AZUL, align=PP_ALIGN.CENTER)

    # --- Diagrama de clases ------------------------------------------------
    s = nuevo(prs, 5)
    titulo(s, "Diagrama de clases del nucleo", size=20, idx=0, algn=PP_ALIGN.LEFT)
    imagen_del_ph(s, 13, os.path.join(DIAG, "diagrama_clases_nucleo.png"))
    lista(s, 14, ["22 clases · 5 paquetes · sin librerias externas   |   "
                  "docs/diagramas/diagrama_clases_nucleo.png"],
          size=12, color=GRIS, marco="", align=PP_ALIGN.LEFT, gap=0)

    # =================================================================
    # BLOQUE 3 - ISA y pruebas tecnicas (min 16-30)
    # =================================================================
    separador(prs, "03", "ISA del procesador y pruebas tecnicas",
              "formatos · instrucciones · ALU · syscall · 201 pruebas     min. 16 – 30")

    # --- Que es el ISA ----------------------------------------------------
    s = nuevo(prs, 1)
    titulo(s, "La ISA es el contrato del procesador")
    lista(s, 10, [
        "**Interfaz** entre el software y el hardware: que puede hacer la CPU",
        "Define tres cosas: **conjunto de instrucciones**, **registros** y "
        "**modos de direccionamiento**",
        "Programa  ->  compilador / ensamblador  ->  **ISA**  ->  microarquitectura  ->  "
        "circuitos",
        "“Quien controla la ISA, controla que se puede construir sobre ella”",
        "MIPS: especificacion **disponible** y formatos fijos  ->  la implementamos completa",
        "En el codigo: `Opcode`, `Funct`, `FormatoInstruccion` (paquete `instruccion`)",
    ], size=15, gap=11)

    # --- 32 bits / formatos -----------------------------------------------
    s = nuevo(prs, 1)
    titulo(s, "Una instruccion = una palabra de 32 bits")
    borrar_ph(s, 10)
    formatos = [
        ("Formato R", [("opcode", 6), ("rs", 5), ("rt", 5), ("rd", 5), ("shamt", 5),
                       ("funct", 6)]),
        ("Formato I", [("opcode", 6), ("rs", 5), ("rt", 5), ("inmediato 16", 16)]),
        ("Formato J", [("opcode", 6), ("direccion 26", 26)]),
    ]
    y = 1.45
    for nombre, campos in formatos:
        etiqueta(s, 0.55, y + 0.14, 1.35, 0.3, nombre, size=12, color=AZUL,
                 align=PP_ALIGN.LEFT)
        x = 1.95
        total = sum(c[1] for c in campos)
        wtotal = 7.4
        for i, (nom, bits) in enumerate(campos):
            w = wtotal * bits / total
            caja(s, x, y, w, 0.58, "", size=9,
                 fill=AZUL if i == 0 else AZUL_CLARO,
                 color=BLANCO, borde=AZUL_MED, shape=MSO_SHAPE.RECTANGLE)
            etiqueta(s, x, y + 0.07, w, 0.24, campos[i][0], size=9.5,
                     color=BLANCO if i == 0 else AZUL, font=None)
            etiqueta(s, x, y + 0.29, w, 0.24, "%d" % bits, size=8.5,
                     color=AZUL_CLARO if i == 0 else GRIS, bold=False)
            x += w
        y += 0.86
    bloque(s, 0.55, 4.05, 8.9, 1.05,
           ["opcode = bits **31..26**  ->  `Decodificador.opcode(palabra)` = "
            "`(palabra >>> 26) & 0x3F`",
            "Ejemplo  `addi $t0, $zero, 5`  ->  **001000 00000 01000 "
            "0000000000000101**",
            "Codigo: `src/instruccion/Decodificador.java` y `Instruccion.codificarR/I/J`"],
           fill=GRIS_CLARO, borde=AZUL, size=12, color=GRIS, align=PP_ALIGN.LEFT, gap=6)

    # --- Instrucciones soportadas ------------------------------------------
    s = nuevo(prs, 1)
    titulo(s, "ISA implementada: 33 instrucciones")
    borrar_ph(s, 10)
    tabla(s, [
        ["Formato", "Opcode", "Instrucciones", "#"],
        ["R", "0x00", "add addu sub subu and or xor nor slt sltu sll srl mult div "
                      "mfhi mflo jr syscall", "18"],
        ["I", "0x04, 0x05, 0x08 – 0x0F, 0x23, 0x2B",
         "beq bne addi addiu slti sltiu andi ori xori lui lw sw", "12"],
        ["J", "0x02, 0x03", "j  jal", "2"],
        ["Emulador", "0x3F", "halt  (operacion de parada propia)", "1"],
    ], 0.55, 1.40, 8.9, 2.6, col_w=[1.25, 2.0, 4.65, 1.0], size=11.5)
    etiqueta(s, 0.55, 4.35, 8.9, 0.8,
             "En formato R el opcode vale 0 y la operacion concreta la decide el campo "
             "**funct** (6 bits).   Lista completa: `docs/instrucciones.md`",
             size=12, color=AZUL_MED, bold=False)

    # --- 0x00 y 0x3F --------------------------------------------------------
    s = nuevo(prs, 3)
    titulo(s, "0x00  ·  la operacion 0 (NOP)", size=18, idx=17, color=AZUL)
    titulo(s, "0x3F  ·  HALT (operacion de parada)", size=18, idx=16, color=ROJO)
    lista(s, 14, [
        "Palabra `0000…0000`  ->  opcode **0x00** (tipo R) + funct **0x00** (`sll`)",
        "El resultado va a `$rd = $0`  ->  **se ignora**  ->  no pasa nada",
        "Toda celda vacia vale 0, asi que **es un NOP gratis**",
        "Por eso la CPU **nunca muere**: siempre hay algo que ejecutar",
        "`Instruccion.PALABRA_NOP = 0`",
    ], size=13, gap=8)
    lista(s, 15, [
        "Instruccion **propia** del emulador (no existe en el MIPS real)",
        "`UnidadControl` -> senal `halt()` -> `Ejecutor` -> `cpu.detener()`",
        "`detenido = true` y el ciclo ya no avanza",
        "Alternativa: **syscall** con `$v0 = 10` (terminar)",
        "Red: `LIMITE_CICLOS_SEGURIDAD = 100000`",
    ], size=13, gap=8)

    # --- ALU -----------------------------------------------------------------
    s = nuevo(prs, 1)
    titulo(s, "ALU  (Paso 2)")
    borrar_ph(s, 10)
    tabla(s, [
        ["Familia", "Operaciones"],
        ["Aritmetica", "ADD  ADDU  SUB  SUBU"],
        ["Logica", "AND  OR  XOR  NOR"],
        ["Comparacion", "SLT  SLTU  (1 si es menor, si no 0)"],
        ["Desplazamiento", "SLL  SRL"],
        ["Especial", "LUI (imm << 16)   ·   NOP"],
    ], 0.55, 1.40, 5.0, 2.5, col_w=[1.7, 3.3], size=11.5)
    lista_tb = [
        "Dos operandos de 32 bits + una operacion -> resultado de 32 bits",
        "**MIPS no tiene flags**: comparar es `slt` y luego `beq` / `bne`",
        "Desbordamiento: `add`/`sub` **detienen** la ejecucion; `addu`/`subu` envuelven",
        "`mult` y `div` guardan el resultado de 64 bits en **HI:LO** "
        "(`mfhi` / `mflo` lo pasan a un registro)",
        "Codigo: `src/alu/ALU.java` y `src/alu/OperacionALU.java`",
    ]
    tb = s.shapes.add_textbox(IN(5.85), IN(1.45), IN(3.6), IN(3.4))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.clear()
    for i, t in enumerate(lista_tb):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        _pconf(p, marL=0.2, indent=-0.2, space_after=10, line=1.12)
        _runs(p, "•  " + t, 12, GRIS)
    bloque(s, 0.55, 4.15, 8.9, 0.6,
           ["Los programas usan **addu / subu / addiu** para que nunca se detengan "
            "por desbordamiento"],
           fill=AZUL_CLARO, borde=AZUL, size=12.5, color=AZUL, align=PP_ALIGN.CENTER)

    # --- syscall -------------------------------------------------------------
    s = nuevo(prs, 1)
    titulo(s, "syscall: la puerta al sistema operativo")
    borrar_ph(s, 10)
    tabla(s, [
        ["$v0", "Servicio", "Argumento"],
        ["1", "imprimir entero", "$a0 = valor"],
        ["4", "imprimir cadena", "$a0 = direccion terminada en 0"],
        ["5", "leer entero", "el resultado se guarda en $v0"],
        ["11", "imprimir caracter", "$a0 = codigo ASCII"],
        ["10", "terminar programa", "-   (detiene la CPU)"],
    ], 0.55, 1.40, 8.9, 2.4, col_w=[1.1, 3.4, 4.4], size=12)
    tb = s.shapes.add_textbox(IN(0.55), IN(4.05), IN(8.9), IN(1.1))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.clear()
    for i, t in enumerate([
        "En un procesador real es la **frontera entre el programa y el SO** "
        "(entrada y salida en modo usuario); aqui es la consola del emulador",
        "El servicio va en `$v0`, los argumentos en `$a0`  ->  `CPU.syscall()` "
        "(`src/cpu/CPU.java`)",
    ]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        _pconf(p, marL=0.24, indent=-0.24, space_after=8, line=1.1)
        _runs(p, "•  " + t, 12.5, GRIS)

    # --- Pruebas ---------------------------------------------------------------
    s = nuevo(prs, 4)
    titulo(s, "Pruebas tecnicas: 201, todas pasan", idx=11)
    borrar_ph(s, 10)
    lista(s, 12, [
        "**117 unitarias**: ALU, registros, decodificador, segmentos, HALT",
        "**44 de integracion**: ciclo completo, ramas, saltos, syscall",
        "**40 de sistema**: los 5 programas con su salida exacta",
        "Verifican tambien que **la CPU nunca muere**",
        "`./probar.sh`  ->  201 pruebas, **0 fallidas**",
    ], size=12, gap=9)
    imagen(s, os.path.join(IMG, "ev_pruebas.png"), 4.35, 1.75, 5.0, 2.6)

    # =================================================================
    # BLOQUE 4 - Codigos implementados y ejecutados (min 30-45)
    # =================================================================
    separador(prs, "04", "Codigos implementados y ejecutados",
              "5 programas de usuario + evidencia de ejecucion     min. 30 – 45")

    # --- Los 5 programas -------------------------------------------------
    s = nuevo(prs, 1)
    titulo(s, "Los 5 programas de usuario")
    borrar_ph(s, 10)
    tabla(s, [
        ["Programa", "Que hace", "Salida real", "Ciclos"],
        ["sumatoria", "suma de 3 a 10", "Sumatoria de 3 a 10 = 52", "68"],
        ["factorial", "7!", "Factorial de 7 = 5040", "63"],
        ["integral", "1²+2²+3²+4²  (rectangulos)", "… desde x=1 hasta x=5 = 30", "56"],
        ["promedio", "vector {4,8,15,16,23}", "Promedio = 13", "123"],
        ["perceptron", "z = Σ x·w + bias, activacion", "z = 3   y = 1", "62"],
    ], 0.55, 1.40, 8.9, 2.7, col_w=[1.6, 3.1, 3.2, 1.0], size=12)
    etiqueta(s, 0.55, 4.4, 8.9, 0.6,
             "Se arman con `instruccion.ProgramaBuilder` (sin escribir binario a mano) "
             "y terminan siempre con `halt()`.", size=12, color=AZUL_MED, bold=False)

    # --- Programas: uno por diapositiva ---------------------------------
    programas = [
        ("Factorial", "7! = 5040 en 63 ciclos", [
            "**Objetivo**: factorial de 7 -> **5040**",
            "Comparacion **sin flags**: `slt` + `bne`",
            "`mult` + `mflo`: el acumulado pasa por **LO**",
            "Bucle con etiqueta resuelto por `ProgramaBuilder`",
            "Termina con `Impresion.terminar(p)` -> **halt**",
            "Codigo: `src/programas/Factorial.java`",
        ],
            "p.cargarConstante(T1, 1);   // acumulado\n"
            "p.cargarConstante(T2, 1);   // contador i\n"
            "\n"
            "p.etiqueta(\"bucle\");\n"
            "p.slt(T3, T0, T2);           // 1 si n < i\n"
            "p.bne(T3, ZERO, \"fin_bucle\");\n"
            "p.mult(T1, T2);              // acumulado * i\n"
            "p.mflo(T1);\n"
            "p.addiu(T2, T2, 1);          // i++\n"
            "p.j(\"bucle\");\n"
            "\n"
            "p.etiqueta(\"fin_bucle\");\n"
            "Impresion.entero(p, T1);\n"
            "Impresion.terminar(p);       // halt"),
        ("Sumatoria", "3 + 4 + … + 10 = 52 en 68 ciclos", [
            "**Objetivo**: sumar de 3 a 10 -> **52**",
            "Los datos viven en la **RAM** y se leen con `lw`",
            "`addu` / `addiu`: nunca provocan desbordamiento",
            "Bucle que sale con `slt` + `bne` al pasar el final",
            "Salida por `syscall` 1 y 4 (entero y cadena)",
            "Codigo: `src/programas/Sumatoria.java`",
        ],
            "int dirInicial = p.palabra(3);   // en RAM\n"
            "int dirFinal   = p.palabra(10);\n"
            "p.lwAbs(T0, dirInicial);\n"
            "p.lwAbs(T1, dirFinal);\n"
            "\n"
            "p.cargarConstante(T2, 0);       // suma\n"
            "p.mover(T3, T0);                // i = inicial\n"
            "p.etiqueta(\"bucle\");\n"
            "p.slt(T4, T1, T3);              // 1 si final < i\n"
            "p.bne(T4, ZERO, \"fin_bucle\");\n"
            "p.addu(T2, T2, T3);             // suma += i\n"
            "p.addiu(T3, T3, 1);             // i++\n"
            "p.j(\"bucle\");"),
        ("Integral", "1²+2²+3²+4² = 30 en 56 ciclos", [
            "**Objetivo**: integral de x² en [1,5] -> **30**",
            "`mult` + `mflo` calculan x·x",
            "Recorre x de abajo hacia arriba con `addiu`",
            "Acumula en `T2` y compara con `slt` / `beq`",
            "El mas corto: 56 ciclos",
            "Codigo: `src/programas/Integral.java`",
        ],
            "p.cargarConstante(T2, 0);    // acumulado\n"
            "p.mover(T3, T0);             // x = limite inferior\n"
            "\n"
            "p.etiqueta(\"bucle\");\n"
            "p.slt(T4, T3, T1);           // 1 si x < superior\n"
            "p.beq(T4, ZERO, \"fin_bucle\");\n"
            "p.mult(T3, T3);              // x * x\n"
            "p.mflo(T4);\n"
            "p.addu(T2, T2, T4);          // suma\n"
            "p.addiu(T3, T3, 1);          // x++\n"
            "p.j(\"bucle\");\n"
            "\n"
            "p.etiqueta(\"fin_bucle\");"),
        ("Promedio", "{4,8,15,16,23} = 13 en 123 ciclos", [
            "**Objetivo**: promedio entero -> **13** (66 / 5)",
            "Vector en la **RAM**, se recorre con `lw`",
            "Dos bucles: imprimir y sumar (el mas largo)",
            "`div` + `mflo` deja el cociente entero en `T7`",
            "Direccion base + indice: `addu T6, T4, T1`",
            "Codigo: `src/programas/Promedio.java`",
        ],
            "int dirVec = p.palabra(4);   // vector en RAM\n"
            "p.palabra(8); p.palabra(15);\n"
            "p.palabra(16); p.palabra(23);\n"
            "\n"
            "p.etiqueta(\"sumar\");\n"
            "p.slt(T5, T1, T0);           // i < n ?\n"
            "p.beq(T5, ZERO, \"fin_suma\");\n"
            "p.addu(T6, T4, T1);          // direccion\n"
            "p.lw(T6, T6, 0);             // lee vector[i]\n"
            "p.addu(T2, T2, T6);          // suma\n"
            "p.addiu(T1, T1, 1);\n"
            "p.j(\"sumar\");\n"
            "\n"
            "p.div(T2, T0);   p.mflo(T7); // promedio"),
        ("Perceptron", "z = 3 -> y = 1 en 62 ciclos", [
            "**Objetivo**: red neuronal minima -> **z = 3**, **y = 1**",
            "Entradas x, pesos w y bias guardados en **RAM**",
            "Producto por elemento con `mult` + `mflo`, suma en `T2`",
            "Activacion escalonada: `slt` + `bne` (si z < 0, y = 0)",
            "Codigo: `src/programas/Perceptron.java`",
        ],
            "int dirX = p.palabra(1);   // entradas\n"
            "p.palabra(1); p.palabra(0);\n"
            "int dirW = p.palabra(2);   // pesos\n"
            "p.palabra(2); p.palabra(-3);\n"
            "\n"
            "p.mult(T5, T4);            // w[i] * x[i]\n"
            "p.mflo(T7);\n"
            "p.addu(T2, T2, T7);        // z += ...\n"
            "\n"
            "p.slt(T3, T2, ZERO);       // z < 0 ?\n"
            "p.bne(T3, ZERO, \"imprimir\");\n"
            "p.cargarConstante(T9, 1);  // y = 1"),
    ]
    for nombre, bajada, puntos, codigo in programas:
        s = nuevo(prs, 4)
        titulo(s, "%s  -  %s" % (nombre, bajada), idx=11, size=19)
        borrar_ph(s, 10)
        lista(s, 12, puntos, size=12.5, gap=7, line=1.1)
        bloque(s, 4.07, 1.61, 5.31, 3.39, codigo.split("\n"), fill=RGBColor(0x1E, 0x22, 0x2B),
               borde=None, size=9.5, color=RGBColor(0xE6, 0xE9, 0xEF),
               align=PP_ALIGN.LEFT, font=CODIGO, gap=1.5, anchor=MSO_ANCHOR.MIDDLE)

    # --- Evidencia ------------------------------------------------------------
    s = nuevo(prs, 5)
    titulo(s, "Evidencia de ejecucion", size=20, idx=0, algn=PP_ALIGN.LEFT)
    imagen_del_ph(s, 13, os.path.join(IMG, "ev_factorial_fin.png"))
    lista(s, 14, ["java -cp out ui.Emulador factorial --estado   ->   "
                  "5040 en 63 ciclos, PC = 25, LO = 5040, RAM con el programa"],
          size=12, color=GRIS, marco="", align=PP_ALIGN.LEFT, gap=0)

    # --- Demo ------------------------------------------------------------------
    s = nuevo(prs, 1)
    titulo(s, "Demo en vivo (3 minutos)")
    ph_demo = s.placeholders[10]
    ph_demo.left, ph_demo.top = IN(1.344), IN(3.10)
    ph_demo.width, ph_demo.height = IN(8.156), IN(1.95)
    bloque(s, 0.55, 1.45, 8.9, 1.35, [
        "./compilar.sh",
        "java -cp out ui.Emulador factorial --traza --estado",
        "./probar.sh          # 201 pruebas, 0 fallidas",
        "./ventana.sh         # ventana grafica",
    ], fill=RGBColor(0x1E, 0x22, 0x2B), borde=None, size=13,
        color=RGBColor(0x8F, 0xD4, 0x8F), align=PP_ALIGN.LEFT, font=CODIGO, gap=4)
    lista(s, 10, [
        "**`--traza`** imprime cada ciclo: PC y codigo de la instruccion en hexadecimal",
        "**`--estado`** vuelca registros, PC, HI/LO y el mapa de memoria al final",
        "**Ventana**: “Ejecutar 1 ciclo” muestra la instruccion, sus senales de "
        "control y el PC siguiente; `Ctrl+Intro` ejecuta un ciclo",
        "Todo el estado queda en `salidas/*_captura.txt` (evidencia de la practica)",
    ], size=14, gap=12)

    # --- Conclusiones ------------------------------------------------------------
    s = nuevo(prs, 1)
    titulo(s, "Conclusiones")
    lista(s, 10, [
        "**MIPS fue buena eleccion**: ISA reducida, 3 formatos fijos y documentacion clara",
        "**Nucleo modular**: agregar una instruccion toca 4 archivos y nada mas",
        "**La CPU nunca muere**: arranca en 0, las celdas vacias son NOP y solo "
        "`halt` o syscall 10 la paran",
        "**201 pruebas + 5 programas** con salida exacta = evidencia de que funciona",
        "Java permite modelar la circuiteria (bits, registros, senales) sin perder claridad",
    ], size=16, gap=13)

    # --- Trabajo futuro ------------------------------------------------------------
    s = nuevo(prs, 1)
    titulo(s, "Trabajo futuro")
    lista(s, 10, [
        "**Pipeline**: las 5 etapas en paralelo y sus riesgos (datos y de control)",
        "**Mas instrucciones**: `lb` / `sb`, `lui` completo, multiplicacion de 64 bits",
        "**Interrupciones y excepciones**: el desbordamiento como excepcion real",
        "**Cache y segmentacion** de la memoria (hoy es un mapa plano de 300 celdas)",
        "**Ensamblador de texto**: archivo `.s` -> binario, con etiquetas y directivas",
    ], size=16, gap=13)

    # --- Gracias --------------------------------------------------------------------
    s = nuevo(prs, 2)
    s.placeholders[0].top, s.placeholders[0].height = IN(1.95), IN(0.75)
    titulo(s, "Gracias", size=36, algn=PP_ALIGN.CENTER, idx=0)
    lista(s, 10, ["¿Preguntas?   ·   Emulador del procesador MIPS en Java"],
          size=15, color=AZUL_MED, marco="", align=PP_ALIGN.CENTER, gap=0)

    prs.save(SALIDA)
    print("Presentacion guardada en %s  (%d diapositivas)" %
          (os.path.relpath(SALIDA, RAIZ), len(prs.slides)))


if __name__ == "__main__":
    main()
